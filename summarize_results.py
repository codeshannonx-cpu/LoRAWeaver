import os
import json
import argparse
from copy import copy
from collections import defaultdict
from numbers import Number
from pathlib import Path
from nltk.translate.bleu_score import sentence_bleu
import numpy as np
import pandas as pd
from openpyxl import load_workbook

STRUCT_TO_TEXT_DOMAIN = "struct to text"
TRANSLATION_DOMAIN = "translation"

ROW_GROUP_NATURAL_LANGUAGE_GENERATION = 'natural_language_generation'
ROW_GROUP_NATURAL_LANGUAGE_UNDERSTANDING = 'natural_language_understanding'

ROW_GROUP_ORDER = [
    ROW_GROUP_NATURAL_LANGUAGE_GENERATION,
    ROW_GROUP_NATURAL_LANGUAGE_UNDERSTANDING,
]

ROW_GROUP_TITLES = {
    ROW_GROUP_NATURAL_LANGUAGE_GENERATION: 'Natural Language Generation',
    ROW_GROUP_NATURAL_LANGUAGE_UNDERSTANDING: 'Natural Language Understanding',
}

# Write task names here to exclude them from task-wise summaries.
TASK_NAMES_TO_SKIP = [

    
]

# Write domain names here to exclude them from summaries.
DOMAIN_NAMES_TO_SKIP = [
    
]

LATEX_PRIORITY_COLUMNS = [
]

ORACLE_DISPLAY_NAME = 'Oracle'
ORACLE_DISPLAY_NAME_ALIASES = {
    'Oracle',
    'Perfect Selection',
    'best_selection',
    'best selection',
}

LATEX_COLUMN_COLORS = {
    
}

RANKING_EXCLUDED_COLUMNS = {
    'Oracle',
}

# Function to calculate BLEU score
def calculate_bleu(references, candidates):
    scores = [sentence_bleu([ref.split()], cand.split()) for ref, cand in zip(references, candidates)]
    return np.round(np.mean(scores) * 100, 1) if scores else 0

# Function to calculate ROUGE score
def calculate_rouge(references, candidates):
    from rouge import Rouge

    rouge = Rouge()
    scores = rouge.get_scores(candidates, references, avg=True)
    rouge_1 = np.round(scores['rouge-1']['f'] * 100, 1)
    rouge_2 = np.round(scores['rouge-2']['f'] * 100, 1)
    rouge_l = np.round(scores['rouge-l']['f'] * 100, 1)
    return rouge_1, rouge_2, rouge_l


def load_bertscore_metric():
    try:
        import evaluate
    except ImportError as exc:
        raise ImportError(
            "BERTScore evaluation requires the 'evaluate' and 'bert-score' packages. "
            "Install the project requirements before using --struct-to-text-score bertscore."
        ) from exc

    return evaluate.load("bertscore")


def calculate_bertscore(references, candidates, bertscore_metric, lang='en'):
    results = bertscore_metric.compute(predictions=candidates, references=references, lang=lang)
    return np.round(float(np.mean(results['f1'])) * 100, 1)

# Function to calculate Exact Match score
def calculate_em(references, candidates):
    references = [ref.split("\n\n")[0] for ref in references]
    em_scores = [1 if cal_correct(ref, cand) else 0 for ref, cand in zip(references, candidates)]
    return np.round(np.mean(em_scores) * 100, 1) if em_scores else 0

def cal_correct(generated_answer, expected_answer):
    is_correct = generated_answer.strip().lower().replace(".", "") == expected_answer.strip().lower().replace(".", "")
    return is_correct

# Function to process a file
def process_file(file_path):
    with open(file_path, 'r') as file:
        data = json.load(file)

    organized_data = defaultdict(lambda: defaultdict(list))
    for entry in data:
        domain = entry['domain']
        task = entry['task']
        organized_data[domain][task].append(entry)
    
    return organized_data


def discover_json_files(folder_path):
    json_files = []
    root_path = Path(folder_path)

    for file_path in sorted(root_path.rglob('*.json')):
        if not file_path.is_file():
            continue

        relative_path = file_path.relative_to(root_path)
        parent_label = relative_path.parent.as_posix() if relative_path.parent != Path('.') else root_path.name
        json_files.append(
            {
                'relative_path': relative_path.as_posix(),
                'display_name': file_path.stem,
                'directory_name': parent_label,
                'full_path': str(file_path),
            }
        )

    return json_files


def escape_latex_text(value):
    escaped_value = str(value)
    escaped_value = escaped_value.replace('\\', r'\textbackslash{}')

    latex_special_chars = {
        '&': r'\&',
        '%': r'\%',
        '$': r'\$',
        '#': r'\#',
        '_': r'\_',
        '{': r'\{',
        '}': r'\}',
        '~': r'\textasciitilde{}',
        '^': r'\textasciicircum{}',
    }

    for special_char, replacement in latex_special_chars.items():
        escaped_value = escaped_value.replace(special_char, replacement)

    return escaped_value


def get_latex_display_name(column_name, file_metadata):
    display_name = file_metadata.get(column_name, {}).get('display_name', Path(column_name).stem)
    if normalize_display_name(display_name) in {
        normalize_display_name(alias)
        for alias in ORACLE_DISPLAY_NAME_ALIASES
    }:
        return ORACLE_DISPLAY_NAME

    return Path(display_name).stem


def normalize_display_name(value):
    return ''.join(char for char in str(value).casefold() if char.isalnum())


def is_oracle_column(column_name, file_metadata):
    display_name = get_latex_display_name(column_name, file_metadata)
    normalized_aliases = {
        normalize_display_name(alias)
        for alias in ORACLE_DISPLAY_NAME_ALIASES
    }
    return normalize_display_name(display_name) in normalized_aliases


def order_latex_value_columns(value_columns, file_metadata):
    priority_map = {name: index for index, name in enumerate(LATEX_PRIORITY_COLUMNS)}
    original_positions = {column: index for index, column in enumerate(value_columns)}

    return sorted(
        value_columns,
        key=lambda column: (
            priority_map.get(get_latex_display_name(column, file_metadata), len(priority_map)),
            original_positions[column],
        ),
    )


def is_ranking_excluded_column(column_name, file_metadata=None):
    if file_metadata is not None:
        display_name = get_latex_display_name(column_name, file_metadata)
    else:
        display_name = Path(str(column_name)).stem

    return display_name in RANKING_EXCLUDED_COLUMNS


def get_summary_row_group(group_name, group_by):
    if group_by != 'domain':
        return None

    normalized_group_name = normalize_display_name(group_name)
    natural_language_generation_domains = {
        normalize_display_name(STRUCT_TO_TEXT_DOMAIN),
        normalize_display_name(TRANSLATION_DOMAIN),
    }
    if normalized_group_name in natural_language_generation_domains:
        return ROW_GROUP_NATURAL_LANGUAGE_GENERATION
    return ROW_GROUP_NATURAL_LANGUAGE_UNDERSTANDING


def build_average_row(group_column_name, json_file_keys, row_label, score_map):
    row = {group_column_name: row_label}
    for file_key in json_file_keys:
        row[file_key] = np.round(np.mean(score_map[file_key]), 1) if score_map[file_key] else 0
    return row


def update_row_level_oracle_normalized_scores(
    row_average_scores,
    oracle_file_key,
    json_file_keys,
    global_score_map,
    grouped_score_map=None,
    row_group=None,
):
    oracle_score = row_average_scores.get(oracle_file_key)
    if oracle_score is None or oracle_score <= 0:
        return

    for file_key in json_file_keys:
        score = row_average_scores.get(file_key)
        if score is None:
            continue

        normalized_score = (score / oracle_score) * 100
        global_score_map[file_key].append(normalized_score)
        if grouped_score_map is not None and row_group is not None:
            grouped_score_map[row_group][file_key].append(normalized_score)


def get_ranked_values(row_values, rankable_columns):
    numeric_values = [
        float(value)
        for value, should_rank in zip(row_values, rankable_columns)
        if should_rank and isinstance(value, Number)
    ]

    if not numeric_values:
        return None, None

    unique_sorted_values = sorted(set(numeric_values), reverse=True)
    highest_value = unique_sorted_values[0]
    second_highest_value = unique_sorted_values[1] if len(unique_sorted_values) > 1 else None
    return highest_value, second_highest_value


def format_ranked_latex_row_values(row_values, value_columns, file_metadata):
    rankable_columns = [
        not is_ranking_excluded_column(column, file_metadata)
        for column in value_columns
    ]
    highest_value, second_highest_value = get_ranked_values(row_values, rankable_columns)

    formatted_values = []
    for value, should_rank in zip(row_values, rankable_columns):
        if not isinstance(value, Number):
            formatted_values.append(escape_latex_text(value))
            continue

        formatted_value = f"{float(value):.1f}"
        if should_rank and highest_value is not None and float(value) == highest_value:
            formatted_values.append(f"\\textbf{{{formatted_value}}}")
        elif should_rank and second_highest_value is not None and float(value) == second_highest_value:
            formatted_values.append(f"\\underline{{{formatted_value}}}")
        else:
            formatted_values.append(formatted_value)

    return formatted_values

# Function to process all files in a folder and aggregate scores by domain/task and metric
def process_folder(
    folder_path,
    group_by='domain',
    skip_domain_names=None,
    struct_to_text_score='rouge',
    bertscore_lang='en',
):
    if group_by not in {'domain', 'task'}:
        raise ValueError("group_by must be either 'domain' or 'task'")

    grouped_metrics = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    json_files = discover_json_files(folder_path)
    bertscore_metric = None
    skip_domain_name_set = {
        domain_name.strip()
        for domain_name in (skip_domain_names or [])
        if domain_name.strip()
    }

    for file_info in json_files:
        file_key = file_info['relative_path']
        domains_data = process_file(file_info['full_path'])

        for domain, tasks_data in domains_data.items():
            if domain in skip_domain_name_set:
                continue
            for task, entries in tasks_data.items():
                group_name = domain if group_by == 'domain' else task
                metric = entries[0]['metric']
                references = [entry['targets'] for entry in entries]
                candidates = [entry['predicted_answer'] for entry in entries]
                sample_count = len(entries)

                use_bertscore = (
                    domain == STRUCT_TO_TEXT_DOMAIN
                    and metric == 'rouge'
                    and struct_to_text_score in {'bertscore', 'both'}
                )

                if use_bertscore:
                    if bertscore_metric is None:
                        bertscore_metric = load_bertscore_metric()
                    score = calculate_bertscore(
                        references,
                        candidates,
                        bertscore_metric,
                        lang=bertscore_lang,
                    )
                    grouped_metrics[group_name]['bertscore'][file_key].append({'score': score, 'count': sample_count, 'task': task})
                    if struct_to_text_score == 'bertscore':
                        continue

                if metric == 'bleu':
                    score = calculate_bleu(references, candidates)
                    grouped_metrics[group_name][metric][file_key].append({'score': score, 'count': sample_count, 'task': task})
                elif metric == 'rouge':
                    rouge_1, rouge_2, rouge_l = calculate_rouge(references, candidates)
                    grouped_metrics[group_name]['rouge-1'][file_key].append({'score': rouge_1, 'count': sample_count, 'task': task})
                    grouped_metrics[group_name]['rouge-2'][file_key].append({'score': rouge_2, 'count': sample_count, 'task': task})
                    grouped_metrics[group_name]['rouge-l'][file_key].append({'score': rouge_l, 'count': sample_count, 'task': task})
                elif metric == 'em':
                    score = calculate_em(references, candidates)
                    grouped_metrics[group_name][metric][file_key].append({'score': score, 'count': sample_count, 'task': task})
    
    return grouped_metrics

# Function to build a summary table with group and metric averages
def build_summary_dataframe(
    data,
    folder_path,
    group_by='domain',
    skip_tasks=0,
    skip_task_names=None,
    include_averages=True,
    categorize_domains=True,
):
    data_list = []
    json_files = discover_json_files(folder_path)
    json_file_keys = [file_info['relative_path'] for file_info in json_files]
    file_metadata = {file_info['relative_path']: file_info for file_info in json_files}
    oracle_file_key = next(
        (
            file_key
            for file_key in json_file_keys
            if is_oracle_column(file_key, file_metadata)
        ),
        None,
    )
    macro_scores = defaultdict(list)
    weighted_score_sums = defaultdict(float)
    weighted_count_sums = defaultdict(int)
    oracle_normalized_scores = defaultdict(list)
    grouped_rows = defaultdict(list)
    grouped_macro_scores = defaultdict(lambda: defaultdict(list))
    grouped_oracle_normalized_scores = defaultdict(lambda: defaultdict(list))
    group_column_name = 'Domain-Metric' if group_by == 'domain' else 'Task-Metric'
    items = sorted(data.items(), key=lambda item: item[0])

    if group_by == 'task':
        skip_task_names = skip_task_names or []
        skip_task_name_set = {task_name.strip() for task_name in skip_task_names if task_name.strip()}
        items = [(group_name, metrics) for group_name, metrics in items if group_name not in skip_task_name_set]

        if skip_tasks > 0:
            items = items[skip_tasks:]

    for group_name, metrics in items:
        for metric, files in metrics.items():
            row_group = get_summary_row_group(group_name, group_by) if categorize_domains else None
            row = {group_column_name: f"{group_name}-{metric}"}
            row_average_scores = {}
            for file_key in json_file_keys:
                file_entries = files[file_key]
                numeric_scores = [entry['score'] for entry in file_entries if isinstance(entry, dict) and isinstance(entry.get('score'), (int, float))]
                sample_counts = [entry['count'] for entry in file_entries if isinstance(entry, dict) and isinstance(entry.get('score'), (int, float))]
                average_score = np.mean(numeric_scores) if numeric_scores else 0
                row[file_key] = np.round(average_score, 1)
                row_average_scores[file_key] = average_score if numeric_scores else None

                if numeric_scores:
                    macro_scores[file_key].append(average_score)
                    if row_group is not None:
                        grouped_macro_scores[row_group][file_key].append(average_score)
                    weighted_score_sums[file_key] += sum(score * count for score, count in zip(numeric_scores, sample_counts))
                    weighted_count_sums[file_key] += sum(sample_counts)

            if oracle_file_key is not None:
                update_row_level_oracle_normalized_scores(
                    row_average_scores,
                    oracle_file_key,
                    json_file_keys,
                    oracle_normalized_scores,
                    grouped_oracle_normalized_scores,
                    row_group,
                )

            if row_group is not None:
                grouped_rows[row_group].append(row)
            else:
                data_list.append(row)

    if group_by == 'domain' and categorize_domains:
        ordered_rows = []
        section_break_indices = []
        for row_group in ROW_GROUP_ORDER:
            rows = grouped_rows[row_group]
            if not rows:
                continue

            if ordered_rows:
                section_break_indices.append(len(ordered_rows))
            ordered_rows.extend(rows)

        if include_averages and ordered_rows:
            section_break_indices.append(len(ordered_rows))
            ordered_rows.append(
                build_average_row(
                    group_column_name,
                    json_file_keys,
                    'Macro Average',
                    macro_scores,
                )
            )
            sample_weighted_average_row = {group_column_name: 'Sample-Weighted Average'}
            for file_key in json_file_keys:
                sample_weighted_average_row[file_key] = (
                    np.round(weighted_score_sums[file_key] / weighted_count_sums[file_key], 1)
                    if weighted_count_sums[file_key]
                    else 0
                )
            ordered_rows.append(sample_weighted_average_row)
            ordered_rows.append(
                build_average_row(
                    group_column_name,
                    json_file_keys,
                    'Oracle-Normalized Average',
                    oracle_normalized_scores,
                )
            )
            ordered_rows.append(
                build_average_row(
                    group_column_name,
                    json_file_keys,
                    'NLU Oracle-Normalized Average',
                    grouped_oracle_normalized_scores[ROW_GROUP_NATURAL_LANGUAGE_UNDERSTANDING],
                )
            )

        data_list = ordered_rows

    if group_by == 'domain' and not categorize_domains and include_averages:
        data_list.append(
            build_average_row(
                group_column_name,
                json_file_keys,
                'Macro Average',
                macro_scores,
            )
        )
        sample_weighted_average_row = {group_column_name: 'Sample-Weighted Average'}
        for file_key in json_file_keys:
            sample_weighted_average_row[file_key] = (
                np.round(weighted_score_sums[file_key] / weighted_count_sums[file_key], 1)
                if weighted_count_sums[file_key]
                else 0
            )
        data_list.append(sample_weighted_average_row)
        data_list.append(
            build_average_row(
                group_column_name,
                json_file_keys,
                'Final Accuracy',
                oracle_normalized_scores,
            )
        )

    df = pd.DataFrame(data_list)
    columns_ordered = [group_column_name] + json_file_keys
    df = df[columns_ordered]
    df.attrs['section_break_indices'] = section_break_indices if group_by == 'domain' and categorize_domains else []

    if group_by == 'domain':
        return df

    if not include_averages:
        return df

    macro_average_row = {group_column_name: 'Macro Average'}
    for file_key in json_file_keys:
        macro_average_row[file_key] = np.round(np.mean(macro_scores[file_key]), 1) if macro_scores[file_key] else 0

    sample_weighted_average_row = {group_column_name: 'Sample-Weighted Average'}
    for file_key in json_file_keys:
        sample_weighted_average_row[file_key] = (
            np.round(weighted_score_sums[file_key] / weighted_count_sums[file_key], 1)
            if weighted_count_sums[file_key]
            else 0
        )

    oracle_normalized_average_row = {group_column_name: 'Oracle-Normalized Average'}
    for file_key in json_file_keys:
        oracle_normalized_average_row[file_key] = (
            np.round(np.mean(oracle_normalized_scores[file_key]), 1)
            if oracle_normalized_scores[file_key]
            else 0
        )

    df = pd.concat(
        [
            df,
            pd.DataFrame([
                macro_average_row,
                sample_weighted_average_row,
                oracle_normalized_average_row,
            ]),
        ],
        ignore_index=True,
    )

    return df

# Function to convert the summary DataFrame to LaTeX format
def convert_to_latex_modified(df, folder_path):
    formatted_df = df.copy()
    section_break_indices = list(df.attrs.get('section_break_indices', []))
    group_column_name = formatted_df.columns[0]
    json_files = discover_json_files(folder_path)
    file_metadata = {file_info['relative_path']: file_info for file_info in json_files}
    value_columns = [column for column in formatted_df.columns if column != group_column_name]
    value_columns = order_latex_value_columns(value_columns, file_metadata)
    formatted_df = formatted_df[[group_column_name] + value_columns]

    formatted_rows = []
    for _, row in formatted_df.iterrows():
        formatted_row = {group_column_name: escape_latex_text(row[group_column_name])}
        ranked_values = format_ranked_latex_row_values(
            [row[column] for column in value_columns],
            value_columns,
            file_metadata,
        )
        for column, formatted_value in zip(value_columns, ranked_values):
            formatted_row[column] = formatted_value
        formatted_rows.append(formatted_row)

    formatted_df = pd.DataFrame(formatted_rows, columns=formatted_df.columns)

    column_specs = ['l']
    for column in value_columns:
        display_name = get_latex_display_name(column, file_metadata)
        color = LATEX_COLUMN_COLORS.get(display_name)
        if color:
            column_specs.append('>{' + '\\columncolor{' + color + '}}c')
        else:
            column_specs.append('c')
    column_format = ''.join(column_specs)
    latex_body = formatted_df.to_latex(index=False, header=False, escape=False, column_format=column_format)
    body_lines = latex_body.splitlines()
    header_start = next(
        (index for index, line in enumerate(body_lines) if '\\midrule' in line),
        next(index for index, line in enumerate(body_lines) if '\\toprule' in line) + 1,
    )

    grouped_headers = []
    current_directory = None
    current_span = 0
    for column in value_columns:
        directory_name = file_metadata.get(column, {}).get('directory_name', Path(column).parent.as_posix() or Path(folder_path).name)
        if directory_name == current_directory:
            current_span += 1
        else:
            if current_directory is not None:
                grouped_headers.append((current_directory, current_span))
            current_directory = directory_name
            current_span = 1
    if current_directory is not None:
        grouped_headers.append((current_directory, current_span))

    top_header_cells = [escape_latex_text(group_column_name)]
    for directory_name, span in grouped_headers:
        top_header_cells.append(f"\\multicolumn{{{span}}}{{c}}{{{escape_latex_text(directory_name)}}}")

    sub_header_cells = ['']
    for column in value_columns:
        display_name = get_latex_display_name(column, file_metadata)
        sub_header_cells.append(escape_latex_text(display_name))

    custom_header_lines = [
        ' & '.join(top_header_cells) + r' \\',
        ' & '.join(sub_header_cells) + r' \\',
        r'\midrule',
    ]

    body_content_lines = body_lines[header_start + 1:]
    if section_break_indices:
        updated_body_content_lines = []
        row_index = 0
        for line in body_content_lines:
            if line.endswith(r'\\'):
                if row_index in section_break_indices:
                    updated_body_content_lines.append(r'\midrule')
                updated_body_content_lines.append(line)
                row_index += 1
                continue

            updated_body_content_lines.append(line)
        body_content_lines = updated_body_content_lines

    final_lines = body_lines[:header_start] + custom_header_lines + body_content_lines
    return '\n'.join(final_lines)

# Function to save the summary DataFrame to Excel
def save_summary_to_excel(df, output_path):
    df.to_excel(output_path, index=False)

    workbook = load_workbook(output_path)
    worksheet = workbook.active
    rankable_columns = {
        cell.column
        for cell in worksheet[1][1:]
        if not is_ranking_excluded_column(cell.value)
    }

    for row in worksheet.iter_rows(min_row=2, min_col=2, max_col=worksheet.max_column):
        numeric_cells = [
            cell
            for cell in row
            if cell.column in rankable_columns and isinstance(cell.value, Number)
        ]
        if not numeric_cells:
            continue

        highest_value, second_highest_value = get_ranked_values(
            [cell.value for cell in numeric_cells],
            [True] * len(numeric_cells),
        )

        for cell in numeric_cells:
            updated_font = copy(cell.font)
            if highest_value is not None and float(cell.value) == highest_value:
                updated_font = copy(updated_font)
                updated_font.bold = True
            elif second_highest_value is not None and float(cell.value) == second_highest_value:
                updated_font = copy(updated_font)
                updated_font.underline = 'single'

            cell.font = updated_font

    workbook.save(output_path)

def parse_args():
    parser = argparse.ArgumentParser(description='Summarize evaluation results by domain or task.')
    parser.add_argument(
        '--folder-path',
        default='',
        help='Folder containing JSON result files.',
    )
    parser.add_argument(
        '--group-by',
        choices=['domain', 'task'],
        default='domain',
        help='Whether to aggregate scores domain-wise or task-wise.',
    )
    parser.add_argument(
        '--output-file',
        default=None,
        help='Optional Excel output filename. Defaults to summary_results_<group>.xlsx.',
    )
    parser.add_argument(
        '--skip-tasks',
        type=int,
        default=0,
        help='Number of alphabetically sorted tasks to skip before summarizing when using task grouping.',
    )
    parser.add_argument(
        '--skip-domains',
        nargs='*',
        default=None,
        help='Optional list of domain names to exclude from summaries.',
    )
    parser.add_argument(
        '--struct-to-text-score',
        choices=['rouge', 'bertscore', 'both'],
        default='rouge',
        help='Metric to use for the struct to text domain. Use both to report ROUGE and BERTScore.',
    )
    parser.add_argument(
        '--bertscore-lang',
        default='en',
        help='Language passed to BERTScore when --struct-to-text-score bertscore is  used.',
    )
    averages_group = parser.add_mutually_exclusive_group()
    averages_group.add_argument(
        '--averages',
        dest='include_averages',
        action='store_true',
        help='Include average rows in the summary table.',
    )
    averages_group.add_argument(
        '--no-averages',
        dest='include_averages',
        action='store_false',
        help='Omit average rows from the summary table.',
    )
    domain_categories_group = parser.add_mutually_exclusive_group()
    domain_categories_group.add_argument(
        '--domain-categories',
        dest='categorize_domains',
        action='store_true',
        help='Split domain summaries into NLG/NLU sections.',
    )
    domain_categories_group.add_argument(
        '--no-domain-categories',
        dest='categorize_domains',
        action='store_false',
        help='Keep all domain rows together and add one final accuracy row.',
    )
    parser.set_defaults(include_averages=True)
    parser.set_defaults(categorize_domains=True)
    return parser.parse_args()


def main():
    args = parse_args()
    folder_path = args.folder_path
    if args.skip_domains is not None:
        skip_domain_names = args.skip_domains
    else:
        skip_domain_names = DOMAIN_NAMES_TO_SKIP
    processed_data = process_folder(
        folder_path,
        group_by=args.group_by,
        skip_domain_names=skip_domain_names,
        struct_to_text_score=args.struct_to_text_score,
        bertscore_lang=args.bertscore_lang,
    )
    summary_df = build_summary_dataframe(
        processed_data,
        folder_path,
        group_by=args.group_by,
        skip_tasks=args.skip_tasks,
        skip_task_names=TASK_NAMES_TO_SKIP,
        include_averages=args.include_averages,
        categorize_domains=args.categorize_domains,
    )
    latex_table = convert_to_latex_modified(summary_df, folder_path)
    output_file = args.output_file or f"summary_results_{args.group_by}.xlsx"
    excel_output_path = os.path.join(folder_path, output_file)
    save_summary_to_excel(summary_df, excel_output_path)

    print(latex_table)
    print(f"Grouping mode: {args.group_by}")
    print(f"Struct-to-text metric: {args.struct_to_text_score}")
    print(f"Average rows: {'included' if args.include_averages else 'omitted'}")
    print(f"Domain categories: {'enabled' if args.categorize_domains else 'disabled'}")
    if skip_domain_names:
        print(f"Skipped domains: {', '.join(skip_domain_names)}")
    print(f"Excel summary saved to: {excel_output_path}")


if __name__ == '__main__':
    main()
